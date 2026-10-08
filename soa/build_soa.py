"""Build the Statement of Applicability workbook for the Saarland University ISMS exercise.

Usage: python3 build_soa.py ../risk-register/saarland-university-risk-register.xlsx out.xlsx [sources.md]

The risk register is read to fill in which risks each Annex A control treats
(column L of the Risk Register sheet), so the two documents stay in step.
"""
import datetime as dt
import re
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

REGISTER, OUT = sys.argv[1], sys.argv[2]

FONT = "Arial"
f = lambda **k: Font(name=FONT, **k)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
CALC_FILL = PatternFill("solid", fgColor="595959")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

THEMES = {"5": "Organisational", "6": "People", "7": "Physical", "8": "Technological"}
STATUSES = ["Implemented", "Partially implemented", "Planned", "Not implemented", "Not applicable"]
REASONS = {"R": "Risk treatment", "L": "Legal / regulatory", "C": "Contractual", "B": "Baseline good practice"}

# Control owner roles (invented, matching the register's roles)
CISO = "Chief Information Security Officer (CISO)"
HIZ = "Head of University IT Centre (HIZ)"
HR = "Director of Human Resources"
DPO = "Data Protection Officer"
FM = "Head of Facilities Management"
VPR = "Vice President for Research"
BOARD = "University Executive Board"
LEGAL = "Head of Legal Affairs"
PROC = "Head of Procurement"

# (control, short paraphrased title, applicable, other reasons, status, owner,
#  justification, evidence / reference document)
# Titles are short paraphrases, not the ISO wording. Risk links (reason R) are
# added automatically from the register.
CONTROLS = [
 ("A.5.1", "Security policy set and reviewed", True, "L,B", "Implemented", BOARD,
  "Top-level policy is the basis of the ISMS; required to show management direction.", "Information Security Policy"),
 ("A.5.2", "Security roles and duties", True, "B", "Partially implemented", BOARD,
  "Responsibilities are split between HIZ, faculties and institutes; roles must be written down.", "ISMS Roles Charter"),
 ("A.5.3", "Separating conflicting duties", True, "B", "Partially implemented", HIZ,
  "Small admin teams often request and approve their own changes.", "Access Control Guideline"),
 ("A.5.4", "Management duties for security", True, "B", "Implemented", BOARD,
  "Deans and heads of department must enforce the policy in their units.", "Information Security Policy"),
 ("A.5.5", "Contact with public authorities", True, "L", "Implemented", CISO,
  "Breach reports to the data protection supervisor and contact with police and BSI after attacks.", "Incident Response Plan"),
 ("A.5.6", "Contact with specialist groups", True, "B", "Implemented", CISO,
  "Exchange with the university CERT community and research network security teams.", "ISMS Roles Charter"),
 ("A.5.7", "Threat intelligence gathering", True, "", "Not implemented", CISO,
  "Universities are targets for ransomware and research espionage; warnings are not collected systematically.", "Planned: Threat Intelligence Procedure"),
 ("A.5.8", "Security in project management", True, "B", "Partially implemented", CISO,
  "IT and research projects start without a security check today.", "Project Management Guideline"),
 ("A.5.9", "Asset inventory", True, "", "Partially implemented", HIZ,
  "Central systems are listed; institute servers are not.", "Configuration Management Database"),
 ("A.5.10", "Rules for acceptable use of assets", True, "B", "Partially implemented", CISO,
  "IT usage regulations exist but do not cover public AI tools.", "IT Usage Regulations"),
 ("A.5.11", "Returning assets on leaving", True, "B", "Partially implemented", HR,
  "Keys and laptops are returned on request only; guest researchers are often missed.", "HR Security Procedure"),
 ("A.5.12", "Information classification", True, "L", "Planned", VPR,
  "Research, health and exam data need different protection; no scheme exists yet.", "Planned: Classification Guideline"),
 ("A.5.13", "Labelling information", True, "B", "Planned", VPR,
  "Follows the classification scheme so staff can see how to handle a document.", "Planned: Classification Guideline"),
 ("A.5.14", "Transferring information", True, "L", "Partially implemented", DPO,
  "Personal data is shared by email, cloud links and with external partners.", "Data Transfer Guideline"),
 ("A.5.15", "Access control rules", True, "L", "Partially implemented", HIZ,
  "Need-to-know rules exist for central systems, not for shared drives.", "Access Control Guideline"),
 ("A.5.16", "Managing identities", True, "", "Partially implemented", HIZ,
  "One identity per person across staff, students and guests is the base of all access.", "Identity Management Concept"),
 ("A.5.17", "Handling authentication secrets", True, "B", "Implemented", HIZ,
  "Password rules and secure reset via the service desk are in place.", "Password Guideline"),
 ("A.5.18", "Granting and reviewing access rights", True, "", "Partially implemented", HIZ,
  "Rights are granted on request but rarely reviewed or removed.", "Access Control Guideline"),
 ("A.5.19", "Security with suppliers", True, "C", "Planned", PROC,
  "Core services (network, cloud, campus management) depend on external suppliers.", "Planned: Supplier Security Guideline"),
 ("A.5.20", "Security terms in supplier contracts", True, "L,C", "Partially implemented", PROC,
  "Data processing agreements are signed; security clauses are not standard yet.", "Contract templates"),
 ("A.5.21", "ICT supply chain security", True, "", "Not implemented", HIZ,
  "Single network uplink and cloud services create supply chain exposure.", "Planned: Supplier Security Guideline"),
 ("A.5.22", "Monitoring supplier services and changes", True, "C", "Not implemented", PROC,
  "Supplier performance and security changes are not reviewed.", "Planned: Supplier Security Guideline"),
 ("A.5.23", "Security for cloud services", True, "L", "Partially implemented", HIZ,
  "Collaboration and teaching tools run in the cloud; selection rules exist, exit plans do not.", "Cloud Usage Guideline"),
 ("A.5.24", "Planning for incident management", True, "L", "Partially implemented", CISO,
  "A response process exists on paper but has not been exercised.", "Incident Response Plan"),
 ("A.5.25", "Assessing security events", True, "B", "Partially implemented", CISO,
  "Service desk forwards suspicious events; triage criteria are informal.", "Incident Response Plan"),
 ("A.5.26", "Responding to incidents", True, "L", "Partially implemented", CISO,
  "Needed to contain attacks and meet the 72-hour GDPR breach reporting deadline.", "Incident Response Plan"),
 ("A.5.27", "Learning from incidents", True, "B", "Not implemented", CISO,
  "No lessons-learned review after incidents so far.", "Incident Response Plan"),
 ("A.5.28", "Collecting evidence", True, "L", "Not implemented", CISO,
  "Evidence may be needed for police, disciplinary or court cases.", "Planned: Forensics Procedure"),
 ("A.5.29", "Security during disruption", True, "", "Planned", CISO,
  "Security must be kept up while systems are restored after a major attack.", "Business Continuity Plan"),
 ("A.5.30", "ICT readiness for continuity", True, "", "Planned", HIZ,
  "Enrolment, exams and payroll depend on IT being restored in time.", "Business Continuity Plan"),
 ("A.5.31", "Legal, regulatory and contract requirements", True, "L,C", "Partially implemented", LEGAL,
  "Data protection law, research funder terms and export rules apply.", "Register of legal requirements"),
 ("A.5.32", "Intellectual property rights", True, "L,C", "Implemented", LEGAL,
  "Software licences and research results must be protected.", "Licence Management Procedure"),
 ("A.5.33", "Protecting records", True, "L", "Implemented", LEGAL,
  "Student, exam and HR records have legal retention periods.", "Records Retention Schedule"),
 ("A.5.34", "Privacy and personal data", True, "L", "Partially implemented", DPO,
  "Personal data of students, staff, applicants and study participants is processed.", "Records of Processing Activities"),
 ("A.5.35", "Independent security review", True, "B", "Planned", BOARD,
  "Internal audit of the ISMS is planned once core controls are in place.", "Planned: Internal Audit Programme"),
 ("A.5.36", "Compliance with security policies and standards", True, "B", "Not implemented", CISO,
  "Faculties and institutes are not checked against the policy.", "Planned: Internal Audit Programme"),
 ("A.5.37", "Documented operating procedures", True, "", "Partially implemented", HIZ,
  "Key systems depend on a few people's knowledge.", "HIZ operations handbook"),

 ("A.6.1", "Background checks", True, "L", "Partially implemented", HR,
  "Limited by German employment and data protection law; applied to administrator roles only.", "HR Security Procedure"),
 ("A.6.2", "Employment terms covering security", True, "L", "Implemented", HR,
  "Contracts refer to confidentiality and the IT usage regulations.", "Employment contract templates"),
 ("A.6.3", "Security awareness and training", True, "", "Partially implemented", CISO,
  "Phishing is the main entry point; training today is occasional emails only.", "Planned: Awareness Programme"),
 ("A.6.4", "Disciplinary process", True, "L", "Implemented", HR,
  "Public service and student regulations already provide a process.", "HR Security Procedure"),
 ("A.6.5", "Duties after leaving or changing role", True, "L", "Partially implemented", HR,
  "Confidentiality continues after people leave; access changes on role moves are missed.", "HR Security Procedure"),
 ("A.6.6", "Confidentiality agreements", True, "L,C", "Implemented", LEGAL,
  "Needed for guests, contractors and industry research partners.", "NDA templates"),
 ("A.6.7", "Remote working", True, "B", "Partially implemented", CISO,
  "Staff work from home and travel to conferences; VPN is available.", "Remote Working Guideline"),
 ("A.6.8", "Reporting security events", True, "L", "Implemented", CISO,
  "Everyone can report via the IT service desk; needed for breach deadlines.", "Incident Response Plan"),

 ("A.7.1", "Physical security perimeters", True, "B", "Partially implemented", FM,
  "Open campus by design; perimeters are set for the data centre and server rooms only.", "Physical Security Guideline"),
 ("A.7.2", "Physical entry controls", True, "B", "Implemented", FM,
  "Data centre and server rooms use badge access.", "Physical Security Guideline"),
 ("A.7.3", "Securing offices and facilities", True, "B", "Partially implemented", FM,
  "Offices handling exams and HR files need locks and closed windows.", "Physical Security Guideline"),
 ("A.7.4", "Physical security monitoring", True, "B", "Partially implemented", FM,
  "Alarm system covers the data centre; institute server rooms are not monitored.", "Physical Security Guideline"),
 ("A.7.5", "Protection from physical and environmental threats", True, "", "Partially implemented", FM,
  "Fire, water and heat could destroy central IT.", "Data Centre Operations Manual"),
 ("A.7.6", "Working in secure areas", True, "B", "Implemented", FM,
  "Visitors and contractors in the data centre are escorted.", "Data Centre Operations Manual"),
 ("A.7.7", "Clear desk and clear screen", True, "B", "Partially implemented", CISO,
  "Shared offices and public areas expose documents and screens.", "IT Usage Regulations"),
 ("A.7.8", "Siting and protecting equipment", True, "B", "Implemented", FM,
  "Servers and network devices sit in locked rooms and racks.", "Data Centre Operations Manual"),
 ("A.7.9", "Assets used off-site", True, "", "Partially implemented", CISO,
  "Laptops are taken home and to conferences.", "Remote Working Guideline"),
 ("A.7.10", "Storage media", True, "L", "Partially implemented", HIZ,
  "USB drives and external disks carry research and personal data.", "Cryptography Guideline"),
 ("A.7.11", "Supporting utilities", True, "", "Implemented", FM,
  "Data centre has UPS and generator; cooling is redundant.", "Data Centre Operations Manual"),
 ("A.7.12", "Cabling security", True, "B", "Implemented", FM,
  "Campus fibre and data centre cabling run in protected ducts.", "Network Security Concept"),
 ("A.7.13", "Equipment maintenance", True, "B", "Implemented", FM,
  "Maintenance contracts for UPS, cooling and core hardware.", "Data Centre Operations Manual"),
 ("A.7.14", "Secure disposal or reuse of equipment", True, "L", "Partially implemented", HIZ,
  "Old disks may hold personal data; wiping is done centrally but not by institutes.", "Disposal Procedure"),

 ("A.8.1", "User endpoint devices", True, "", "Partially implemented", CISO,
  "Many laptops are not centrally managed or encrypted.", "Endpoint Security Guideline"),
 ("A.8.2", "Privileged access", True, "", "Partially implemented", HIZ,
  "Admin accounts have broad standing rights.", "Access Control Guideline"),
 ("A.8.3", "Restricting access to information", True, "L", "Partially implemented", HIZ,
  "Shared drives and cloud folders are often open to large groups.", "Access Control Guideline"),
 ("A.8.4", "Access to source code", True, "B", "Partially implemented", HIZ,
  "In-house applications are kept in a central Git service; rights are not reviewed.", "Secure Development Guideline"),
 ("A.8.5", "Secure authentication", True, "", "Partially implemented", HIZ,
  "Single sign-on protects many services; MFA is being rolled out.", "Identity Management Concept"),
 ("A.8.6", "Capacity management", True, "", "Partially implemented", HIZ,
  "Peak loads at enrolment and exam registration and from attacks.", "Capacity planning notes"),
 ("A.8.7", "Malware protection", True, "", "Partially implemented", HIZ,
  "Antivirus on endpoints; servers lack detection and response tooling.", "Endpoint Security Guideline"),
 ("A.8.8", "Technical vulnerability management", True, "", "Partially implemented", CISO,
  "Public IP ranges are scanned yearly; institute systems are patched irregularly.", "Patch and Vulnerability Procedure"),
 ("A.8.9", "Configuration management", True, "B", "Partially implemented", HIZ,
  "Central servers follow baselines; institute systems do not.", "Configuration Management Database"),
 ("A.8.10", "Deleting information", True, "L", "Partially implemented", DPO,
  "Personal data must be deleted when no longer needed.", "Records Retention Schedule"),
 ("A.8.11", "Data masking", True, "L", "Planned", VPR,
  "Research data with health information needs pseudonymisation.", "Planned: Pseudonymisation Guideline"),
 ("A.8.12", "Data leakage prevention", True, "L", "Not implemented", CISO,
  "Personal and research data can leave by email, cloud and AI tools.", "Planned: DLP rules for email"),
 ("A.8.13", "Backup", True, "", "Partially implemented", HIZ,
  "Nightly backups exist; no offline copy and restores are not tested.", "Backup Concept"),
 ("A.8.14", "Redundancy", True, "", "Partially implemented", HIZ,
  "Campus management and core services need to survive a single failure.", "Business Continuity Plan"),
 ("A.8.15", "Logging", True, "L", "Partially implemented", HIZ,
  "Logs are kept per system, not centrally; admin actions are not logged everywhere.", "Logging and Monitoring Concept"),
 ("A.8.16", "Monitoring activities", True, "", "Partially implemented", CISO,
  "Attacks are found late; no central security monitoring.", "Logging and Monitoring Concept"),
 ("A.8.17", "Clock synchronisation", True, "B", "Implemented", HIZ,
  "Central time servers keep log timestamps consistent.", "Network Security Concept"),
 ("A.8.18", "Use of privileged utility programs", True, "B", "Partially implemented", HIZ,
  "Admin tools are restricted on central systems only.", "Access Control Guideline"),
 ("A.8.19", "Installing software on operational systems", True, "B", "Partially implemented", HIZ,
  "Central systems use change control; institutes install freely.", "Change Management Procedure"),
 ("A.8.20", "Network security", True, "", "Implemented", HIZ,
  "Edge firewall and upstream DDoS filtering; remaining DoS risk is accepted.", "Network Security Concept"),
 ("A.8.21", "Security of network services", True, "C", "Implemented", HIZ,
  "Uplink, eduroam and VPN have defined service and security terms.", "Network Security Concept"),
 ("A.8.22", "Network segregation", True, "", "Partially implemented", HIZ,
  "Lab equipment and building systems share the campus network.", "Network Security Concept"),
 ("A.8.23", "Web filtering", True, "B", "Partially implemented", HIZ,
  "Limited to blocking known malicious sites; no content filtering, out of respect for academic freedom.", "Network Security Concept"),
 ("A.8.24", "Cryptography", True, "L", "Partially implemented", CISO,
  "Encryption protects laptops, research data and data in transit.", "Cryptography Guideline"),
 ("A.8.25", "Secure development life cycle", True, "", "Planned", CISO,
  "HIZ and institutes build their own web applications.", "Planned: Secure Development Guideline"),
 ("A.8.26", "Application security requirements", True, "B", "Not implemented", CISO,
  "Security needs are not written into specifications or tenders.", "Planned: Secure Development Guideline"),
 ("A.8.27", "Secure system architecture and engineering", True, "B", "Partially implemented", HIZ,
  "Design principles are applied to central systems only.", "Network Security Concept"),
 ("A.8.28", "Secure coding", True, "", "Planned", CISO,
  "In-house code has had injection and access flaws.", "Planned: Secure Development Guideline"),
 ("A.8.29", "Security testing in development and acceptance", True, "", "Planned", CISO,
  "Applications go live without security testing.", "Planned: Secure Development Guideline"),
 ("A.8.30", "Outsourced development", False, "", "Not applicable", CISO,
  "Excluded: within the ISMS scope no software development is outsourced. Bought standard software is covered by A.5.19 to A.5.22. Review if this changes.", "Scope statement"),
 ("A.8.31", "Separating development, test and production", True, "B", "Partially implemented", HIZ,
  "Central applications have test systems; institute applications often do not.", "Change Management Procedure"),
 ("A.8.32", "Change management", True, "B", "Partially implemented", HIZ,
  "A change board covers central systems.", "Change Management Procedure"),
 ("A.8.33", "Test information", True, "L", "Partially implemented", HIZ,
  "Real student data is sometimes copied into test systems.", "Change Management Procedure"),
 ("A.8.34", "Protecting systems during audit testing", True, "B", "Implemented", CISO,
  "Scans and penetration tests are agreed with system owners beforehand.", "Patch and Vulnerability Procedure"),
]
assert len(CONTROLS) == 93, len(CONTROLS)
assert len({c[0] for c in CONTROLS}) == 93

# ---------------- risks per control, from the register ----------------
reg = load_workbook(REGISTER)["Risk Register"]
RISKS_FOR = {}
for row in reg.iter_rows(min_row=5, values_only=True):
    rid, annex = row[0], row[11]
    if not rid or not annex:
        continue
    for ctl in re.findall(r"A\.\d+\.\d+", annex):
        RISKS_FOR.setdefault(ctl, []).append(rid)
unknown = set(RISKS_FOR) - {c[0] for c in CONTROLS}
assert not unknown, f"register cites controls not in Annex A: {unknown}"

wb = Workbook()

# ---------------- Read Me ----------------
rm = wb.active
rm.title = "Read Me"
rm.column_dimensions["A"].width = 110
linked = sum(1 for c in CONTROLS if c[0] in RISKS_FOR)
lines = [
 ("Statement of Applicability: Saarland University", f(bold=True, size=14, color="1F3864")),
 ("PORTFOLIO EXERCISE. Not an official document of Saarland University.", f(bold=True, color="C00000")),
 ("", None),
 ("Purpose", f(bold=True)),
 ("Statement of Applicability (SoA) as required by ISO/IEC 27001:2022 clause 6.1.3 d). It lists all 93 controls of Annex A, says whether each is applicable, why it is included or excluded, and how far it is implemented.", f()),
 ("Control titles are short paraphrases, not the wording of the standard. The full text is in ISO/IEC 27001:2022 and the guidance in ISO/IEC 27002:2022 (paid, not included).", f()),
 ("All statuses, owners, justifications and documents are invented for illustration and do not describe the university's real security posture.", f()),
 ("", None),
 ("Scope (assumed)", f(bold=True)),
 ("Information and IT services of the central university administration and of the University IT Centre (HIZ) at the Saarbrücken and Homburg campuses, including services offered to faculties and institutes. Systems run by institutes themselves are in scope where they connect to the campus network.", f()),
 ("", None),
 ("Sheets", f(bold=True)),
 ("SoA: one row per Annex A control with applicability, reasons for inclusion, justification, implementation status, owner, evidence and the risks it treats.", f()),
 (f"Related risks are read from the risk register (column L of its Risk Register sheet) when the workbook is built: {linked} controls treat at least one of its risks.", f()),
 ("Summary: counts by theme and status, calculated with formulas.", f()),
 ("Lists: values for the drop-down menus.", f()),
 ("Sources: standards and other references used.", f()),
 ("", None),
 ("Reasons for inclusion", f(bold=True)),
 *[(f"{k} = {v}", f()) for k, v in REASONS.items()],
 ("", None),
 ("How to edit", f(bold=True)),
 ("Light yellow cells are inputs. Applicable, Status and the reason flags use drop-downs. An excluded control must say why in the justification; its status should be Not applicable.", f()),
 ("", None),
 (f"Version 0.1, prepared {dt.date(2026,10,8):%d %B %Y}.", f(italic=True, color="595959")),
]
for i, (text, font) in enumerate(lines, 1):
    c = rm.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        c.font = font

# ---------------- Lists ----------------
ls = wb.create_sheet("Lists")
lists = {"A": ("Applicable", ["Yes", "No"]), "B": ("Status", STATUSES), "C": ("Flag", ["X"])}
for col, (h, vals) in lists.items():
    ls[f"{col}1"] = h; ls[f"{col}1"].font = f(bold=True)
    for i, v in enumerate(vals, 2):
        ls[f"{col}{i}"] = v; ls[f"{col}{i}"].font = f()
    ls.column_dimensions[col].width = 22

# ---------------- SoA ----------------
ws = wb.create_sheet("SoA", 1)
ws["A1"] = "Statement of Applicability: ISO/IEC 27001:2022 Annex A"
ws["A1"].font = f(bold=True, size=14, color="1F3864")
ws["A2"] = "Portfolio exercise for Saarland University. Paraphrased control titles; all statuses and details are invented. Not an official university document."
ws["A2"].font = f(italic=True, color="C00000")

cols = [
 ("Control", 9, "in"), ("Control (paraphrased title)", 34, "in"), ("Theme", 15, "in"),
 ("Applicable", 11, "in"),
 ("R: Risk treatment", 9, "in"), ("L: Legal / regulatory", 10, "in"), ("C: Contractual", 10, "in"), ("B: Baseline", 9, "in"),
 ("Justification for inclusion or exclusion", 55, "in"), ("Implementation status", 15, "in"),
 ("Control owner (role)", 28, "in"), ("Evidence / reference document", 30, "in"),
 ("Related risks (from register)", 16, "in"), ("Number of related risks", 10, "calc"),
]
HR = 4
for j, (name, w, kind) in enumerate(cols, 1):
    c = ws.cell(row=HR, column=j, value=name)
    c.font = f(bold=True, color="FFFFFF")
    c.fill = CALC_FILL if kind == "calc" else HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = BORDER
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[HR].height = 45

first = HR + 1
for i, (ctl, title, appl, reasons, status, owner, just, doc) in enumerate(CONTROLS):
    row = first + i
    risks = RISKS_FOR.get(ctl, [])
    flags = set(filter(None, reasons.split(",")))
    if risks:
        flags.add("R")
    if appl and not flags:
        raise ValueError(f"{ctl} is applicable but has no reason")
    values = {
        1: ctl, 2: title, 3: THEMES[ctl.split(".")[1]], 4: "Yes" if appl else "No",
        5: "X" if "R" in flags else None, 6: "X" if "L" in flags else None,
        7: "X" if "C" in flags else None, 8: "X" if "B" in flags else None,
        9: just, 10: status, 11: owner, 12: doc,
        13: ", ".join(risks) or None,
        14: f'=IF(M{row}="",0,LEN(M{row})-LEN(SUBSTITUTE(M{row},",",""))+1)',
    }
    for j, v in values.items():
        c = ws.cell(row=row, column=j, value=v)
        c.font = f()
        c.border = BORDER
        c.alignment = CENTER if j in (1, 3, 4, 5, 6, 7, 8, 10, 13, 14) else WRAP
        if cols[j - 1][2] == "in":
            c.fill = INPUT_FILL
last = first + len(CONTROLS) - 1

ws.freeze_panes = ws.cell(row=first, column=3)
ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(cols))}{last}"

dv_appl = DataValidation(type="list", formula1="=Lists!$A$2:$A$3", allow_blank=False)
dv_status = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${len(STATUSES)+1}", allow_blank=False)
dv_flag = DataValidation(type="list", formula1="=Lists!$C$2:$C$2", allow_blank=True)
for dv in (dv_appl, dv_status, dv_flag):
    ws.add_data_validation(dv)
dv_appl.add(f"D{first}:D{last}")
dv_status.add(f"J{first}:J{last}")
dv_flag.add(f"E{first}:H{last}")

STATUS_COLORS = {"Implemented": "C6EFCE", "Partially implemented": "FFE699", "Planned": "F4B183",
                 "Not implemented": "F8CBAD", "Not applicable": "D9D9D9"}
for name, bg in STATUS_COLORS.items():
    ws.conditional_formatting.add(f"J{first}:J{last}",
        CellIsRule(operator="equal", formula=[f'"{name}"'], fill=PatternFill("solid", fgColor=bg),
                   font=Font(name=FONT, bold=True)))
ws.conditional_formatting.add(f"D{first}:D{last}",
    CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="D9D9D9"),
               font=Font(name=FONT, bold=True, color="C00000")))

# ---------------- Summary ----------------
su = wb.create_sheet("Summary", 2)
su["A1"] = "SoA Summary"; su["A1"].font = f(bold=True, size=14, color="1F3864")
su["A2"] = f"Calculated from the SoA sheet (rows {first} to {last}). Portfolio exercise."; su["A2"].font = f(italic=True)
S = "SoA!"
rng = lambda c: f"{S}${c}${first}:${c}${last}"

def header(row, col, names):
    for k, h in enumerate(names):
        c = su.cell(row=row, column=col + k, value=h)
        c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def put(row, col, value, bold=False, fmt=None):
    c = su.cell(row=row, column=col, value=value)
    c.font = f(bold=bold); c.border = BORDER
    if col > 1:
        c.alignment = Alignment(horizontal="center")
    if fmt:
        c.number_format = fmt
    return c

kpis = [
 ("Annex A controls", f'=COUNTA({rng("A")})', None),
 ("Applicable", f'=COUNTIF({rng("D")},"Yes")', None),
 ("Excluded", f'=COUNTIF({rng("D")},"No")', None),
 ("Applicable and fully implemented", f'=COUNTIFS({rng("D")},"Yes",{rng("J")},"Implemented")', None),
 ("Implementation rate (fully implemented / applicable)", "=IFERROR(B8/B6,0)", "0%"),
 ("Applicable controls linked to at least one risk", f'=COUNTIFS({rng("D")},"Yes",{rng("N")},">0")', None),
]
header(4, 1, ["KPI", "Value"])
for i, (k, formula, fmt) in enumerate(kpis, 5):
    put(i, 1, k); put(i, 2, formula, bold=True, fmt=fmt)

top = 13
su.cell(row=top - 1, column=1, value="Applicable controls by theme and implementation status").font = f(bold=True)
status_cols = STATUSES[:-1]
header(top, 1, ["Theme"] + status_cols + ["Applicable", "Excluded"])
for i, theme in enumerate(THEMES.values(), top + 1):
    put(i, 1, theme, bold=True)
    for k, st in enumerate(status_cols, 2):
        put(i, k, f'=COUNTIFS({rng("C")},$A{i},{rng("D")},"Yes",{rng("J")},{get_column_letter(k)}${top})')
    put(i, len(status_cols) + 2, f'=COUNTIFS({rng("C")},$A{i},{rng("D")},"Yes")', bold=True)
    put(i, len(status_cols) + 3, f'=COUNTIFS({rng("C")},$A{i},{rng("D")},"No")')
tot = top + len(THEMES) + 1
put(tot, 1, "Total", bold=True)
for k in range(2, len(status_cols) + 4):
    L = get_column_letter(k)
    put(tot, k, f"=SUM({L}{top+1}:{L}{tot-1})", bold=True)
for k, st in enumerate(status_cols, 2):
    su.cell(row=top, column=k).fill = PatternFill("solid", fgColor=STATUS_COLORS[st])
    su.cell(row=top, column=k).font = f(bold=True)

rtop = tot + 3
su.cell(row=rtop - 1, column=1, value="Reasons for inclusion (a control can have several)").font = f(bold=True)
header(rtop, 1, ["Reason", "Controls"])
for i, (code, col) in enumerate(zip(REASONS, "EFGH"), rtop + 1):
    put(i, 1, f"{code}: {REASONS[code]}")
    put(i, 2, f'=COUNTIF({rng(col)},"X")')

su.column_dimensions["A"].width = 46
for col in "BCDEFG":
    su.column_dimensions[col].width = 14
su.row_dimensions[top].height = 32

# ---------------- Sources ----------------
# (section, source, url, what it supports, used in, access)
SOURCES = [
 ("Standard", "ISO/IEC 27001:2022 Information security management systems: Requirements (with Amd 1:2024)",
  "https://www.iso.org/standard/27001",
  "Clause 6.1.3 d) requires a Statement of Applicability; Annex A lists the 93 controls in four themes",
  "Whole SoA; control numbers and themes", "Paid standard (about CHF 155), no copy included"),
 ("Standard", "ISO/IEC 27002:2022 Information security controls", "https://www.iso.org/standard/75652.html",
  "Purpose and guidance for each control, used to write the paraphrased titles and justifications",
  "Columns B and I", "Paid standard, no copy included"),
 ("Law", "EU General Data Protection Regulation (GDPR), Regulation (EU) 2016/679",
  "https://eur-lex.europa.eu/eli/reg/2016/679/oj",
  "Art. 32 security of processing; Art. 33 breach notification within 72 hours",
  "Reason L; A.5.26, A.5.34, A.8.10, A.8.11 and others", "Free"),
 ("Law", "Basic Law for the Federal Republic of Germany (Grundgesetz), Art. 5(3)",
  "https://www.gesetze-im-internet.de/gg/art_5.html",
  "Freedom of science, research and teaching, the reason web filtering is limited to known malicious sites",
  "A.8.23", "Free"),
 ("Internal", "Saarland University risk register (this project)", "../risk-register/saarland-university-risk-register.xlsx",
  "Risks R01 to R18 and the Annex A controls that treat them", "Column M (related risks), reason R", "Included in this repository"),
]
so = wb.create_sheet("Sources")
so["A1"] = "Sources"; so["A1"].font = f(bold=True, size=14, color="1F3864")
so["A2"] = "Checked 8 October 2026. Everything not listed here (statuses, owners, documents, scope) is invented for the exercise."
so["A2"].font = f(italic=True)
shdr = ["Type", "Source", "Link", "What it supports", "Used in", "Access"]
for j, h in enumerate(shdr, 1):
    c = so.cell(row=4, column=j, value=h); c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
for i, src in enumerate(SOURCES, 5):
    for j, v in enumerate(src, 1):
        c = so.cell(row=i, column=j, value=v); c.font = f(); c.alignment = WRAP; c.border = BORDER
        if j == 3 and v.startswith("http"):
            c.hyperlink = v
            c.font = f(color="0563C1", underline="single")
for col, w in zip("ABCDEF", (12, 40, 45, 55, 28, 24)):
    so.column_dimensions[col].width = w

if len(sys.argv) > 3:
    with open(sys.argv[3], "w") as md:
        md.write("# Sources for the Saarland University Statement of Applicability\n\n")
        md.write("Portfolio exercise, not an official university document. Sources checked on 8 October 2026. "
                 "Control titles are paraphrased; all statuses, owners, documents and the scope are invented.\n\n")
        for sec in dict.fromkeys(s[0] for s in SOURCES):
            md.write(f"## {sec}\n\n")
            for _, name, url, what, used, access in (s for s in SOURCES if s[0] == sec):
                md.write(f"- [{name}]({url})  \n  Supports: {what}.  \n  Used in: {used}. Access: {access}.\n")
            md.write("\n")
        md.write("## Files in this folder\n\n"
                 "- `saarland-university-soa.xlsx`: the SoA (sheets Read Me, SoA, Summary, Lists, Sources).\n"
                 "- `build_soa.py`: Python script that generates the workbook and this file "
                 "(`python3 build_soa.py ../risk-register/saarland-university-risk-register.xlsx out.xlsx sources.md`).\n"
                 "- `sources.md`: this list.\n\n"
                 "The ISO standards are paid and cannot be shared. The GDPR and the Basic Law are free to read at the links above.\n")

wb.save(OUT)
print("saved", OUT, len(CONTROLS), "controls,", linked, "linked to risks")
