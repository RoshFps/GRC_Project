import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]

FONT = "Arial"
f = lambda **k: Font(name=FONT, **k)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CALC_FILL = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
DISCLAIMER = ("Portfolio exercise based on ISO 22301 and ISO/IEC 27001:2022 (A.5.29, A.5.30). "
              "Public facts only; all internal details are invented. Not an official university document.")

PROCESS = "Exam registration and grade recording"
SYSTEM = "LSF exam administration (HIS-Portal)"

# Impact scale per category, levels 1 to 5
SCALE = [
 ("Students & teaching",
  ["No noticeable effect", "Some students delayed by hours; workaround via exam office",
   "Many students miss a registration or see no grades for a few days",
   "Students in several faculties cannot sit planned exams or miss deadlines",
   "Exams for a whole term must be postponed; progression, BAfoeG or visa status of students affected"]),
 ("Legal & regulatory",
  ["None", "Minor deviation from internal procedure",
   "Deadlines in exam regulations must be extended by one exam board",
   "Several exam boards must change deadlines; appeals expected; possible GDPR breach",
   "Exam results legally contestable; reportable GDPR breach (Art. 33) with lost records"]),
 ("Reputation",
  ["None", "Complaints to the service desk", "Student council and social media criticism",
   "Regional press coverage", "National press; ministry asks for a report"]),
 ("Operations & workload",
  ["Normal workload", "Exam offices absorb extra calls", "Overtime in exam offices; manual lists needed",
   "Manual registration across faculties; backlog of several days",
   "Backlog of weeks; staff from other units must help"]),
 ("Financial",
  ["Under EUR 5k", "EUR 5k to 25k", "EUR 25k to 100k", "EUR 100k to 500k", "Over EUR 500k"]),
]
CATS = [c for c, _ in SCALE]

TIMES = ["4 hours", "1 day", "2 days", "3 days", "1 week", "2 weeks", "4 weeks"]
HOURS = [4, 24, 48, 72, 168, 336, 672]

# Impact scores per category over time (1 to 5), invented assessment
PEAK = {
 "Students & teaching":   [2, 3, 4, 4, 5, 5, 5],
 "Legal & regulatory":    [1, 2, 3, 4, 5, 5, 5],
 "Reputation":            [1, 2, 3, 3, 4, 5, 5],
 "Operations & workload": [2, 3, 3, 4, 4, 5, 5],
 "Financial":             [1, 1, 1, 2, 2, 3, 3],
}
OFFPEAK = {
 "Students & teaching":   [1, 1, 1, 2, 2, 3, 4],
 "Legal & regulatory":    [1, 1, 1, 1, 2, 2, 3],
 "Reputation":            [1, 1, 1, 1, 2, 2, 3],
 "Operations & workload": [1, 1, 2, 2, 3, 3, 4],
 "Financial":             [1, 1, 1, 1, 1, 1, 2],
}
PEAK_NOTES = {
 "Students & teaching": "Registration windows close on fixed dates and late registration is not possible where it is mandatory. After 2 days the last-minute registrations (assumed 40% of the total) are blocked.",
 "Legal & regulatory": "Deadlines are set in the exam regulations; only exam boards can extend them. After 3 days several boards must decide in parallel and equal treatment of students is hard to show.",
 "Reputation": "Exam registration affects nearly every student at the same time, so complaints spread fast.",
 "Operations & workload": "Exam offices fall back to written registration. Each day of outage adds several thousand forms to type in later (assumption).",
 "Financial": "Mainly overtime, vendor emergency support and possible extra exam dates.",
}
OFFPEAK_NOTES = {
 "Students & teaching": "Outside registration and grading weeks few students need the system on a given day; grade viewing can wait.",
 "Legal & regulatory": "No deadline falls due in most off-peak weeks.",
 "Reputation": "Limited visibility until the outage becomes long.",
 "Operations & workload": "Exam offices can queue certificate and transcript requests for a while.",
 "Financial": "Low.",
}

# Process steps
STEPS = [
 ("1", "Set up exams", "Exam offices create exam dates, rooms and registration deadlines in LSF for each module.", "Exam offices (Prüfungssekretariate)", "Before each registration window"),
 ("2", "Student registration", "Students log in with their university account and register online; each registration is confirmed with an iTAN. The deadline is shown at each exam.", "Students", "Registration window (peak)"),
 ("3", "Withdrawal", "Students withdraw from exams online within the withdrawal period.", "Students", "Until shortly before the exam"),
 ("4", "Attendance lists", "Exam offices export lists of registered students for examiners and room planning.", "Exam offices", "Days before each exam"),
 ("5", "Grade entry", "Examiners or exam offices enter grades in LSF.", "Examiners, exam offices", "Grading deadline (peak)"),
 ("6", "Results published", "Students see results and registration status in the HIS-Portal.", "Students", "After grade entry"),
 ("7", "Transcripts and certificates", "Exam offices issue Transcripts of Records and degree certificates from the recorded grades.", "Exam offices", "All year"),
]

# Dependencies: (name, type, used in steps, current recovery capability, needed by (h), linked risks, note)
DEPS = [
 ("LSF application servers", "IT system", "1-7", "Rebuild on new virtual machines and restore from backup; estimated 72 h (assumption)", 24, "R07, R01, R12", "Core of the process. Today no standby system."),
 ("Exam database (registrations and grades)", "Data", "2-7", "Nightly backup on disk in the same building; up to 24 h data loss", 24, "R01, R12, R08", "Holds legally relevant records."),
 ("Single sign-on and directory service", "IT system", "2, 5, 6", "Restore from backup; estimated 24 h (assumption)", 4, "R01, R08, R09", "Without login nobody can register or enter grades."),
 ("iTAN confirmation", "IT function", "2, 3", "Part of LSF; no separate recovery", 24, "R07", "Fallback: written registration at the exam office."),
 ("Campus network and external uplink", "Infrastructure", "2, 3, 5, 6", "Single uplink path; provider SLA (assumption)", 24, "R05, R13", "Most students register from home."),
 ("Data centre power and cooling", "Facility", "All", "UPS for short outages; generator test planned", 4, "R12", "Same building as the backups."),
 ("Email and university website", "IT system", "Communication", "Hosted separately from LSF (assumption)", 4, "R05, R13, R02", "Needed to tell students about extensions."),
 ("LSF administrators at HIZ (2 people)", "People", "All", "Informal deputy only", 4, "R17", "Key-person risk: thin runbooks."),
 ("Exam office staff in each faculty", "People", "1, 2, 4, 7", "Can work on paper for a limited time", 4, "R17", "Carry the manual fallback."),
 ("Software vendor support", "Supplier", "All", "Standard support contract, next business day (assumption)", 24, "R05, R07", "Needed for database repair and upgrades."),
]

# Linked risks from the register: (id, title, how it disrupts the process, inherent score, BIA note)
LINKED = [
 ("R07", "Campus management system failure during enrolment or exam registration", "Direct outage of LSF during a registration window.", 8,
  "Main scenario of this BIA. The planned manual fallback and deadline extension procedure (due 28 Feb 2027) is written up in the continuity plan."),
 ("R01", "Ransomware on central IT infrastructure", "Encrypts LSF servers, database and possibly the nightly backups.", 20,
  "Worst case for RTO and RPO: backups in the same environment may be lost. Offline backup (R01 action) is a precondition for the RPO target."),
 ("R12", "Fire, water or power failure in the data centre", "Destroys servers and backups kept in the same building.", 10,
  "Second backup site (R12 action, due 31 Oct 2026) is needed for any recovery."),
 ("R05", "Outage of a critical IT or network supplier", "Uplink loss stops students registering from home; vendor delay slows repair.", 8,
  "Fallback: on-campus PC pools and written registration."),
 ("R13", "Denial-of-service attack on website and network", "Overload during the registration deadline.", 9,
  "Accepted risk in the register; the BIA shows it is most harmful in peak weeks. Suggest review of the acceptance."),
 ("R15", "Attacks detected too late", "A compromise of LSF is found late, so clean restore points are older.", 16,
  "Late detection widens the real data loss beyond the RPO."),
 ("R17", "Loss of key IT staff and knowledge", "Only two LSF administrators know the restore steps.", 9,
  "Runbook for LSF restore is part of the actions below."),
 ("R08", "Misuse of privileged administrator accounts", "Grades or registrations changed without trace.", 10,
  "Integrity issue: restore from backup must be checked against change logs."),
]

# Actions to close BIA gaps: (id, action, closes gap, owner, due, linked risk)
ACTIONS = [
 ("BIA-A1", "Replicate the exam database to a second site with transaction logs every 15 minutes", "RPO 24 h -> 1 h", "Head of University IT Centre (HIZ)", "30 Apr 2027", "R01, R12"),
 ("BIA-A2", "Keep a warm standby LSF server at the second site; test switchover each term", "RTO 72 h -> 24 h", "Head of University IT Centre (HIZ)", "30 Jun 2027", "R07, R12"),
 ("BIA-A3", "Write and test an LSF restore runbook with a named deputy", "Key-person dependency", "Head of University IT Centre (HIZ)", "31 Mar 2027", "R17"),
 ("BIA-A4", "Agree a standard deadline extension rule with all exam board chairs in advance", "Legal impact at peak", "Vice President for Teaching", "28 Feb 2027", "R07"),
 ("BIA-A5", "Prepare paper registration forms and an import template for each faculty exam office", "Manual workaround", "Head of Student Services", "28 Feb 2027", "R07"),
 ("BIA-A6", "Upgrade vendor support to 4-hour response during registration and grading weeks", "Supplier recovery time", "Head of University IT Centre (HIZ)", "31 Dec 2026", "R05, R07"),
 ("BIA-A7", "Run the tabletop exercise in this folder before the WiSe 2026/27 registration window", "Untested plan", "Head of Student Services", "15 Dec 2026", "R07"),
]

SOURCES = [
 ("Saarland University: Fragen und Antworten A bis Z (ZPL)", "https://www.uni-saarland.de/en/einrichtung/zpl/studienorganisation/abisz.html",
  "Exam registration and results run in LSF; deadlines shown per exam; late registration not possible where registration is mandatory."),
 ("Saarland University, Systems Engineering: Prüfungsanmeldung und -termine", "https://www.uni-saarland.de/fachrichtung/systems-engineering/studium/pruefungsanmeldung-und-termine.html",
  "Online registration with university account and iTAN; results in the HIS-Portal; if online registration fails, register in writing at the exam office within the deadline."),
 ("Saarland University: Semestertermine Studienjahr 2026/2027 (PDF)", "https://www.uni-saarland.de/fileadmin/upload/dezernat/ls/semestertermine/Studienjahr2026-2027.pdf",
  "WiSe 2026/27 lecture period 12 Oct 2026 to 5 Feb 2027; SoSe 2027 lecture period 5 Apr to 16 Jul 2027."),
 ("ISO 22301:2019 Business continuity management systems", "https://www.iso.org/standard/75106.html",
  "BIA, MTPD, RTO and RPO terms (clause 8.2.2). Paid standard, no copy included."),
 ("BSI-Standard 200-4 Business Continuity Management", "https://www.bsi.bund.de/DE/Themen/Unternehmen-und-Organisationen/Standards-und-Zertifizierung/IT-Grundschutz/BSI-Standards/BSI-Standard-200-4-Business-Continuity-Management/bsi-standard-200-4_Business_Continuity_Management_node.html",
  "German BCM method incl. BIA; used as background for the approach."),
 ("ISO/IEC 27001:2022 Annex A", "https://www.iso.org/standard/27001",
  "A.5.29 Information security during disruption; A.5.30 ICT readiness for business continuity."),
]


def header(ws, row, labels, widths=None):
    for i, h in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=DISCLAIMER):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=14, color="1F3864")
    ws["A2"] = sub
    ws["A2"].font = f(italic=True, size=9, color="595959")


def cell(ws, r, c, v, fill=None, bold=False, align=WRAP):
    x = ws.cell(row=r, column=c, value=v)
    x.font = f(bold=bold)
    x.alignment = align
    x.border = BORDER
    if fill:
        x.fill = fill
    return x


wb = Workbook()

# Read Me
ws = wb.active
ws.title = "Read Me"
title(ws, "Business Impact Analysis: " + PROCESS)
rows = [
 ("Organisation", "Saarland University (Saarbrücken and Homburg campuses)"),
 ("Process", PROCESS + ", supported by " + SYSTEM),
 ("Process owner", "Head of Student Services (same owner as risk R07)"),
 ("Why this process", "Risk R07 in the register: the campus management system failing during exam registration. Nearly every student depends on it within short, fixed windows, and deadlines come from exam regulations, so impact rises fast."),
 ("Result", "Peak MTPD 2 days, RTO 24 hours, RPO 1 hour. Off-peak MTPD 4 weeks, RTO 1 week, RPO 24 hours. Today's capability (about 72 h restore, 24 h data loss) misses the peak targets; see sheet Actions."),
 ("Method", "Impact rated 1 to 5 per category at seven points in time, for peak and off-peak weeks. MTPD is the first point where the highest impact reaches the threshold on sheet Impact Over Time (default 4 = Major). RTO must be at most half of MTPD to leave time for the backlog."),
 ("Sheets", "Process, Impact Over Time, RTO & RPO, Dependencies, Linked Risks, Actions, Impact Scale, Sources"),
 ("Colour key", "Yellow cells are inputs; blue cells are calculated."),
 ("Related files", "continuity-plan.md (one-page plan), tabletop-exercise.md (exercise scenario), ../risk-register/ (risk register)"),
 ("Status", "Draft for review, October 2026"),
]
for i, (k, v) in enumerate(rows, 4):
    cell(ws, i, 1, k, bold=True)
    cell(ws, i, 2, v)
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 110

# Process
ws = wb.create_sheet("Process")
title(ws, "Process overview")
header(ws, 4, ["Step", "Name", "What happens", "Who", "When"], [7, 26, 70, 30, 30])
for i, s in enumerate(STEPS, 5):
    for j, v in enumerate(s, 1):
        cell(ws, i, j, v, align=CENTER if j == 1 else WRAP)
r = 5 + len(STEPS) + 1
cell(ws, r, 1, "Peak periods (assumed for WiSe 2026/27)", bold=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
peaks = [
 ("Registration window", "4 to 22 January 2027 (assumption); the last 3 days carry about 40% of registrations (assumption)"),
 ("Exam period", "From 8 February 2027, after the lecture period ends on 5 February 2027 (public date)"),
 ("Grading deadline", "Mid-March 2027 (assumption), before the semester ends on 31 March 2027 (public date)"),
 ("Volume", "About 60,000 exam registrations per winter term (assumption, based on about 16,300 students)"),
]
for i, (k, v) in enumerate(peaks, r + 1):
    cell(ws, i, 1, k, bold=True)
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    cell(ws, i, 3, v)
    ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)

# Impact Over Time
ws = wb.create_sheet("Impact Over Time")
title(ws, "Impact over time if the process stops")
ws["A3"] = "Unacceptable impact threshold (1 to 5):"
ws["A3"].font = f(bold=True)
ws["D3"] = 4
ws["D3"].fill = INPUT_FILL
ws["D3"].font = f(bold=True)
ws["D3"].border = BORDER
ws["E3"] = "MTPD = first point in time where the highest impact reaches this level."
ws["E3"].font = f(italic=True, size=9)
ws.column_dimensions["A"].width = 26
for col in range(2, 9):
    ws.column_dimensions[get_column_letter(col)].width = 11
ws.column_dimensions["I"].width = 80

refs = {}


def impact_block(top, name, data, notes):
    cell(ws, top, 1, name, bold=True)
    ws.merge_cells(start_row=top, start_column=1, end_row=top, end_column=9)
    ws.cell(row=top, column=1).font = f(bold=True, size=12, color="1F3864")
    header(ws, top + 1, ["Impact category"] + TIMES + ["Reasoning"])
    cell(ws, top + 2, 1, "Hours", bold=True)
    for j, h in enumerate(HOURS, 2):
        cell(ws, top + 2, j, h, align=CENTER)
    first = top + 3
    for i, cat in enumerate(CATS):
        rr = first + i
        cell(ws, rr, 1, cat, bold=True)
        for j, v in enumerate(data[cat], 2):
            cell(ws, rr, j, v, fill=INPUT_FILL, align=CENTER)
        cell(ws, rr, 9, notes[cat])
    last = first + len(CATS) - 1
    mx, un = last + 1, last + 2
    cell(ws, mx, 1, "Highest impact", bold=True)
    cell(ws, un, 1, "Unacceptable?", bold=True)
    for j in range(2, 9):
        L = get_column_letter(j)
        cell(ws, mx, j, f"=MAX({L}{first}:{L}{last})", fill=CALC_FILL, bold=True, align=CENTER)
        cell(ws, un, j, f'=IF({L}{mx}>=$D$3,"Yes","No")', fill=CALC_FILL, align=CENTER)
    m = un + 1
    cell(ws, m, 1, "MTPD", bold=True)
    cell(ws, m, 2, f'=IFERROR(INDEX($B${top+1}:$H${top+1},MATCH("Yes",B{un}:H{un},0)),"> 4 weeks")', fill=CALC_FILL, bold=True, align=CENTER)
    ws.merge_cells(start_row=m, start_column=2, end_row=m, end_column=3)
    cell(ws, m + 1, 1, "MTPD in hours", bold=True)
    cell(ws, m + 1, 2, f'=IFERROR(INDEX($B${top+2}:$H${top+2},MATCH("Yes",B{un}:H{un},0)),9999)', fill=CALC_FILL, align=CENTER)
    for rr in range(first, last + 2):
        ws.conditional_formatting.add(f"B{rr}:H{rr}", CellIsRule(operator="greaterThanOrEqual", formula=["$D$3"],
                                      fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(name=FONT, bold=True, color="C00000")))
    return f"'Impact Over Time'!B{m}", f"'Impact Over Time'!B{m+1}", m + 3


refs["peak_mtpd"], refs["peak_h"], nxt = impact_block(5, "Peak: registration window and grading deadline weeks", PEAK, PEAK_NOTES)
refs["off_mtpd"], refs["off_h"], _ = impact_block(nxt, "Off-peak: all other weeks", OFFPEAK, OFFPEAK_NOTES)

# RTO & RPO
ws = wb.create_sheet("RTO & RPO")
title(ws, "Recovery targets and current capability")
header(ws, 4, ["Measure", "Peak", "Off-peak", "Reasoning"], [34, 16, 16, 100])
rows = [
 ("MTPD (from Impact Over Time)", "=" + refs["peak_mtpd"], "=" + refs["off_mtpd"], "Maximum tolerable period of disruption: after this the damage is no longer acceptable.", CALC_FILL),
 ("MTPD in hours", "=" + refs["peak_h"], "=" + refs["off_h"], "", CALC_FILL),
 ("RTO target (hours)", 24, 168, "Peak: back within one day, which leaves the second day to process paper registrations and announce any extension before the MTPD is reached. Off-peak: one week, so a normal rebuild with vendor support is enough.", INPUT_FILL),
 ("RTO check (RTO at most half of MTPD)", '=IF(B7<=B6/2,"OK","Too long")', '=IF(C7<=C6/2,"OK","Too long")', "Half of MTPD is this kit's rule of thumb to leave time for the backlog after recovery.", CALC_FILL),
 ("RPO target (hours of data loss)", 1, 24, "Peak: registrations are confirmed by students with an iTAN and cannot be repeated after the deadline. On the last day of a window about 10,000 registrations arrive, up to 1,000 in the busiest hour (assumption); a lost day could not be reconstructed. Off-peak: few changes per day, and grades come from examiners' own lists, so they can be entered again.", INPUT_FILL),
 ("Current recovery time (hours)", 72, 72, "Rebuild on new virtual machines, restore the nightly backup, vendor support next business day (assumption).", INPUT_FILL),
 ("Current data loss (hours)", 24, 24, "Daily database backup (existing control in R07); worst case is a failure just before the next backup.", INPUT_FILL),
 ("RTO gap", '=IF(B10>B7,"Gap: "&(B10-B7)&" h","OK")', '=IF(C10>C7,"Gap: "&(C10-C7)&" h","OK")', "Closed by actions BIA-A2, BIA-A3 and BIA-A6.", CALC_FILL),
 ("RPO gap", '=IF(B11>B9,"Gap: "&(B11-B9)&" h","OK")', '=IF(C11>C9,"Gap: "&(C11-C9)&" h","OK")', "Closed by action BIA-A1. Depends on offline backups from R01 and the second site from R12.", CALC_FILL),
 ("Minimum service level during disruption", "Written registration at exam offices; deadline extension", "Queue requests", "From the public Systems Engineering exam page: if online registration fails, students register in writing within the deadline. Details in continuity-plan.md.", None),
]
for i, (k, p, o, why, fill) in enumerate(rows, 5):
    cell(ws, i, 1, k, bold=True)
    cell(ws, i, 2, p, fill=fill, align=CENTER, bold=fill is CALC_FILL)
    cell(ws, i, 3, o, fill=fill, align=CENTER, bold=fill is CALC_FILL)
    cell(ws, i, 4, why)
for rng in ("B8:C8", "B12:C13"):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0]},2)<>"OK"'],
                                  fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(name=FONT, bold=True, color="C00000")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0]},2)="OK"'],
                                  fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=FONT, bold=True, color="006100")))

# Dependencies
ws = wb.create_sheet("Dependencies")
title(ws, "Resources the process depends on")
header(ws, 4, ["Dependency", "Type", "Process steps", "Current recovery capability", "Needed within (hours, peak)", "Linked risks", "Note"],
       [34, 14, 13, 50, 14, 15, 45])
for i, d in enumerate(DEPS, 5):
    for j, v in enumerate(d, 1):
        cell(ws, i, j, v, fill=INPUT_FILL if j == 5 else None, align=CENTER if j in (2, 3, 5, 6) else WRAP)

# Linked Risks
ws = wb.create_sheet("Linked Risks")
title(ws, "Risks from the register that can stop this process")
header(ws, 4, ["Risk ID", "Risk title", "How it disrupts the process", "Inherent score (register)", "What the BIA adds"], [9, 40, 45, 13, 70])
for i, d in enumerate(LINKED, 5):
    for j, v in enumerate(d, 1):
        cell(ws, i, j, v, align=CENTER if j in (1, 4) else WRAP)
ws.cell(row=6 + len(LINKED), column=1, value="Scores copied from ../risk-register/saarland-university-risk-register.xlsx on 8 October 2026.").font = f(italic=True, size=9)

# Actions
ws = wb.create_sheet("Actions")
title(ws, "Actions to close the gaps")
header(ws, 4, ["ID", "Action", "Gap it closes", "Owner (role)", "Due", "Linked risks", "Status"], [9, 70, 26, 34, 13, 13, 12])
for i, a in enumerate(ACTIONS, 5):
    for j, v in enumerate(list(a) + ["Proposed"], 1):
        cell(ws, i, j, v, fill=INPUT_FILL if j == 7 else None, align=CENTER if j in (1, 5, 6, 7) else WRAP)

# Impact Scale
ws = wb.create_sheet("Impact Scale")
title(ws, "Impact scale used in this BIA (aligned with the register's 1 to 5 scale)")
header(ws, 4, ["Category", "1 Negligible", "2 Minor", "3 Moderate", "4 Major", "5 Severe"], [24, 26, 30, 34, 36, 40])
for i, (cat, lv) in enumerate(SCALE, 5):
    cell(ws, i, 1, cat, bold=True)
    for j, v in enumerate(lv, 2):
        cell(ws, i, j, v)

# Sources
ws = wb.create_sheet("Sources")
title(ws, "Sources (checked 8 October 2026)")
header(ws, 4, ["Source", "Link", "What it supports"], [50, 60, 80])
for i, (n, u, s) in enumerate(SOURCES, 5):
    cell(ws, i, 1, n)
    c = cell(ws, i, 2, u)
    c.hyperlink = u
    c.font = f(color="0563C1", underline="single")
    cell(ws, i, 3, s)

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.freeze_panes = "A5" if s.title not in ("Read Me", "Impact Over Time") else None
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToWidth = 1
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.fitToHeight = 0

wb.save(OUT)
print("saved", OUT)
