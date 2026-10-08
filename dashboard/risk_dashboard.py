"""Risk heatmap and KPI dashboard for the Saarland University risk register.

Reads the Risk Register sheet of the workbook and writes:
  heatmap.png, movement.png, owners.png   charts for slides or reports
  risk-dashboard.html                      a single static page with KPIs, the
                                           same charts (as inline SVG) and tables

Usage:
  python3 risk_dashboard.py ../risk-register/saarland-university-risk-register.xlsx
  python3 risk_dashboard.py REGISTER.xlsx --out output --as-of 2026-10-08

Scores, ratings and overdue flags are formulas in the workbook, so they are
recalculated here from likelihood, impact, status and due date. The rating
thresholds are read from the Scoring Method sheet so both stay in step.
"""

import argparse
import html
import io
import logging
import re
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Rectangle
from openpyxl import load_workbook

logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)

# Paper-and-ink look of a printed audit report rather than a web app.
PAPER = "#F5F2EB"
INK = "#1F2124"
MUTED = "#6E695F"
RULE = "#D6CFC0"

# Severity is ordinal, so it gets one warm ramp from sand to oxblood.
RATING_COLOURS = {
    "Low": "#CFD3B0",
    "Medium": "#E8B04F",
    "High": "#CB6234",
    "Critical": "#8B1E2D",
}
RATING_TEXT = {"Low": INK, "Medium": INK, "High": "#FFFFFF", "Critical": "#FFFFFF"}

# Status is a different question from severity, so it uses unrelated hues.
STATUS_COLOURS = {
    "Open": "#2E5AA6",
    "Planned": "#5FA3DE",
    "In progress": "#D9962A",
    "Closed": "#3B9152",
    "Accepted": "#9A68C6",
}
DONE_STATUSES = {"Closed", "Accepted"}

SERIF = "Source Serif 4"
SANS = "IBM Plex Sans"
MONO = "IBM Plex Mono"
GOOGLE_FONTS_CSS = (
    "https://fonts.googleapis.com/css2?family=Source+Serif+4:opsz,wght@8..60,400;8..60,600"
    "&family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&display=swap"
)
FONT_DIR = Path(__file__).with_name(".fonts")


@dataclass
class Risk:
    id: str
    title: str
    category: str
    owner: str
    likelihood: int
    impact: int
    status: str
    due: date | None
    residual_likelihood: int
    residual_impact: int
    remediation: str

    @property
    def inherent(self):
        return self.likelihood * self.impact

    @property
    def residual(self):
        return self.residual_likelihood * self.residual_impact

    @property
    def is_open(self):
        return self.status not in DONE_STATUSES


# ---------------------------------------------------------------- reading

def read_register(path):
    wb = load_workbook(path, data_only=False)
    sheet = wb["Risk Register"]
    risks = []
    for row in sheet.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            break
        due = row[15]
        risks.append(Risk(
            id=row[0], title=row[1], category=row[4], owner=row[5],
            likelihood=int(row[6]), impact=int(row[7]),
            status=row[14], due=due.date() if isinstance(due, datetime) else due,
            residual_likelihood=int(row[18]), residual_impact=int(row[19]),
            remediation=row[13] or "",
        ))

    method = wb["Scoring Method"]
    likelihood_names = {r[0]: r[1] for r in method.iter_rows(min_row=5, max_row=9, values_only=True)}
    impact_names = {r[0]: r[3] for r in method.iter_rows(min_row=5, max_row=9, values_only=True)}
    bands = [(r[0], r[1]) for r in method.iter_rows(min_row=13, max_row=16, max_col=2, values_only=True)]
    return risks, bands, likelihood_names, impact_names


def rating(score, bands):
    name = bands[0][0]
    for band_name, minimum in bands:
        if score >= minimum:
            name = band_name
    return name


def days_overdue(risk, today):
    if not risk.is_open or risk.due is None:
        return 0
    return max(0, (today - risk.due).days)


def compute_kpis(risks, bands, today):
    high = dict(bands)["High"]
    open_risks = [r for r in risks if r.is_open]
    overdue = sorted((r for r in risks if days_overdue(r, today)), key=lambda r: -days_overdue(r, today))
    inherent_total = sum(r.inherent for r in risks)
    residual_total = sum(r.residual for r in risks)
    return {
        "total": len(risks),
        "open": len(open_risks),
        "overdue": overdue,
        "inherent_high": sum(r.inherent >= high for r in risks),
        "residual_high": sum(r.residual >= high for r in risks),
        "reduction": 1 - residual_total / inherent_total if inherent_total else 0,
        "appetite": high,
        "due_soon": [r for r in open_risks if r.due and 0 <= (r.due - today).days <= 60],
    }


# ---------------------------------------------------------------- fonts

def load_fonts():
    """Use Source Serif 4 and IBM Plex for the PNGs, fetched once from Google Fonts.

    Without network access the charts fall back to DejaVu Sans; nothing breaks.
    """
    FONT_DIR.mkdir(exist_ok=True)
    if not any(FONT_DIR.glob("*.ttf")):
        try:
            css = urllib.request.urlopen(GOOGLE_FONTS_CSS, timeout=10).read().decode()
            for i, url in enumerate(re.findall(r"url\((https://[^)]+\.ttf)\)", css)):
                urllib.request.urlretrieve(url, FONT_DIR / f"font-{i}.ttf")
        except OSError:
            pass
    for ttf in FONT_DIR.glob("*.ttf"):
        font_manager.fontManager.addfont(str(ttf))

    plt.rcParams.update({
        "font.family": [SANS, "DejaVu Sans"],
        "font.size": 10,
        "text.color": INK,
        "axes.labelcolor": MUTED,
        "xtick.color": MUTED,
        "ytick.color": MUTED,
        "figure.facecolor": PAPER,
        "axes.facecolor": PAPER,
        "savefig.facecolor": PAPER,
        "svg.fonttype": "none",  # keep text as text so the HTML page styles it
    })


def short_owner(owner):
    """'Head of University IT Centre (HIZ)' reads fine; long role names get trimmed."""
    return owner.replace("Chief Information Security Officer", "CISO").replace(" (CISO)", "")


# ---------------------------------------------------------------- charts

def shorten(text, limit):
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"

def draw_heatmap(risks, bands, likelihood_names, impact_names, tips):
    fig, axes = plt.subplots(1, 2, figsize=(12.4, 6.1))
    fig.subplots_adjust(left=0.12, right=0.985, top=0.86, bottom=0.14, wspace=0.16)
    panels = [
        ("Before treatment", "inherent", lambda r: (r.likelihood, r.impact)),
        ("After planned treatment", "residual", lambda r: (r.residual_likelihood, r.residual_impact)),
    ]
    for ax, (heading, kind, position) in zip(axes, panels):
        cells = defaultdict(list)
        for r in risks:
            cells[position(r)].append(r)

        for likelihood in range(1, 6):
            for impact in range(1, 6):
                band = rating(likelihood * impact, bands)
                here = cells.get((likelihood, impact), [])
                cell = Rectangle((impact - 0.5, likelihood - 0.5), 1, 1,
                                 facecolor=RATING_COLOURS[band], alpha=1 if here else 0.17,
                                 edgecolor=PAPER, linewidth=3)
                ax.add_patch(cell)
                if not here:
                    continue
                tip_id = f"tip-{len(tips)}"
                cell.set_gid(tip_id)
                tips[tip_id] = "\n".join(f"{r.id}  {r.title}" for r in here)
                ids = [r.id for r in here]
                lines = [" ".join(ids[i:i + 3]) for i in range(0, len(ids), 3)]
                ax.text(impact, likelihood, "\n".join(lines), ha="center", va="center",
                        family=[MONO, "DejaVu Sans Mono"], fontsize=8.2, linespacing=1.5,
                        color=RATING_TEXT[band])

        ax.set_xlim(0.5, 5.5)
        ax.set_ylim(0.5, 5.5)
        ax.set_aspect("equal")
        ax.set_xticks(range(1, 6), [f"{i}\n{impact_names[i]}" for i in range(1, 6)], fontsize=8.5)
        ax.set_yticks(range(1, 6), [f"{likelihood_names[i]}  {i}" for i in range(1, 6)], fontsize=8.5)
        ax.tick_params(length=0, pad=6)
        for spine in ax.spines.values():
            spine.set_visible(False)
        ax.set_xlabel("Impact", fontsize=9, labelpad=8)
        if kind == "inherent":
            ax.set_ylabel("Likelihood", fontsize=9, labelpad=8)
        else:
            ax.set_yticklabels([])
        ax.set_title(heading, loc="left", family=[SERIF, "DejaVu Serif"], fontsize=13,
                     color=INK, pad=12)

    handles = [Rectangle((0, 0), 1, 1, facecolor=RATING_COLOURS[name]) for name, _ in bands]
    labels = []
    for i, (name, minimum) in enumerate(bands):
        upper = bands[i + 1][1] - 1 if i + 1 < len(bands) else 25
        labels.append(f"{name}  {minimum}–{upper}")
    fig.legend(handles, labels, loc="upper left", bbox_to_anchor=(0.12, 0.985), ncol=4,
               frameon=False, fontsize=8.5, handlelength=1.1, handleheight=1.1, columnspacing=1.8)
    return fig


def draw_movement(risks, bands, appetite):
    ordered = sorted(risks, key=lambda r: (r.inherent, r.inherent - r.residual), reverse=True)
    fig, ax = plt.subplots(figsize=(12.4, 0.36 * len(ordered) + 1.3))
    fig.subplots_adjust(left=0.36, right=0.975, top=0.93, bottom=0.08)

    for i, (name, minimum) in enumerate(bands):
        upper = bands[i + 1][1] if i + 1 < len(bands) else 25.5
        ax.axvspan(minimum - 0.5, upper - 0.5, color=RATING_COLOURS[name], alpha=0.14, lw=0)
        ax.text((minimum + upper) / 2 - 0.5, -1.0, name.upper(), ha="center", va="bottom",
                fontsize=7.5, color=MUTED, fontweight=500)
    ax.axvline(appetite - 0.5, color=INK, lw=1, ls=(0, (3, 3)))
    ax.text(appetite - 0.25, -0.62, "risk appetite", fontsize=8, color=INK, style="italic", va="center")

    for y, r in enumerate(ordered):
        ax.plot([r.residual, r.inherent], [y, y], color=MUTED, lw=1.4, solid_capstyle="round", zorder=2)
        ax.scatter(r.inherent, y, s=70, facecolor=PAPER, zorder=3, linewidth=2,
                   edgecolor=RATING_COLOURS[rating(r.inherent, bands)])
        ax.scatter(r.residual, y, s=70, zorder=4, edgecolor=MUTED, linewidth=0.8,
                   color=RATING_COLOURS[rating(r.residual, bands)])

    ax.set_yticks(range(len(ordered)), [f"{r.id}   {shorten(r.title, 52)}" for r in ordered],
                  fontsize=8.6, color=INK)
    ax.set_ylim(len(ordered) - 0.4, -1.2)
    ax.set_xlim(0, 25.5)
    ax.set_xticks([1, 5, 10, 15, 20, 25])
    ax.tick_params(length=0, labelsize=8.5)
    ax.set_xlabel("Score (likelihood × impact)", fontsize=9)
    for spine in ax.spines.values():
        spine.set_visible(False)

    ring = plt.Line2D([], [], marker="o", ls="", markersize=8, markerfacecolor=PAPER,
                      markeredgecolor=MUTED, markeredgewidth=2)
    dot = plt.Line2D([], [], marker="o", ls="", markersize=8, color=MUTED)
    fig.legend([ring, dot], ["Inherent (before treatment)", "Residual (after planned treatment)"],
               loc="upper left", bbox_to_anchor=(0.36, 1.0), ncol=2, frameon=False, fontsize=8.5)
    return fig


def draw_owners(risks):
    by_owner = defaultdict(Counter)
    for r in risks:
        by_owner[short_owner(r.owner)][r.status] += 1
    owners = sorted(by_owner, key=lambda o: (sum(by_owner[o].values()), o))
    statuses = [s for s in STATUS_COLOURS if any(c[s] for c in by_owner.values())]

    fig, ax = plt.subplots(figsize=(12.4, 0.5 * len(owners) + 0.9))
    fig.subplots_adjust(left=0.27, right=0.975, top=0.86, bottom=0.04)
    for y, owner in enumerate(owners):
        left = 0
        for status in statuses:
            count = by_owner[owner][status]
            if not count:
                continue
            ax.barh(y, count, left=left, height=0.62, color=STATUS_COLOURS[status],
                    edgecolor=PAPER, linewidth=2)
            text_colour = INK if status in ("Planned", "In progress") else "#FFFFFF"
            ax.text(left + count / 2, y, str(count), ha="center", va="center", fontsize=8.5,
                    color=text_colour, fontweight=500)
            left += count
        ax.text(left + 0.12, y, f"{left}", va="center", fontsize=8.5, color=MUTED)

    ax.set_yticks(range(len(owners)), owners, fontsize=8.8, color=INK)
    ax.set_xticks([])
    ax.set_ylim(-0.5, len(owners) - 0.5)
    ax.tick_params(length=0)
    ax.set_xlim(0, max(sum(c.values()) for c in by_owner.values()) + 0.6)
    for spine in ax.spines.values():
        spine.set_visible(False)
    handles = [Rectangle((0, 0), 1, 1, facecolor=STATUS_COLOURS[s]) for s in statuses]
    fig.legend(handles, statuses, loc="upper left", bbox_to_anchor=(0.27, 0.99), ncol=len(statuses),
               frameon=False, fontsize=8.5, handlelength=1.1, handleheight=1.1, columnspacing=1.8)
    return fig


def save(fig, out_dir, name, tips=None):
    fig.savefig(out_dir / f"{name}.png", dpi=200)
    buffer = io.StringIO()
    fig.savefig(buffer, format="svg")
    plt.close(fig)
    svg = buffer.getvalue()
    svg = svg[svg.index("<svg"):]
    svg = re.sub(r'<svg([^>]*?) width="[^"]+" height="[^"]+"', r'<svg\1', svg, count=1)
    for tip_id, text in (tips or {}).items():
        svg = svg.replace(f'<g id="{tip_id}">', f'<g id="{tip_id}"><title>{html.escape(text)}</title>')
    return svg


# ---------------------------------------------------------------- page

def build_page(risks, bands, kpis, today, source_name, svgs):
    esc = html.escape
    lede = (
        f"{kpis['total']} risks on the register, {kpis['open']} still open. "
        f"{len(kpis['overdue'])} remediation{'s are' if len(kpis['overdue']) != 1 else ' is'} past due. "
        f"Planned treatment brings {kpis['inherent_high']} high or critical risks down to "
        f"{kpis['residual_high']} and cuts the total score by {kpis['reduction']:.0%}."
    )

    def figure(value, label, note="", alert=False):
        cls = ' class="alert"' if alert else ""
        return (f'<div class="figure"><div class="value"{cls}>{value}</div>'
                f'<div class="label">{label}</div><div class="note">{note}</div></div>')

    figures = "".join([
        figure(kpis["total"], "Risks on the register", f"{len({r.category for r in risks})} categories"),
        figure(kpis["open"], "Open risks", "not closed or accepted"),
        figure(len(kpis["overdue"]), "Overdue remediations", "past their due date", alert=bool(kpis["overdue"])),
        figure(f'{kpis["inherent_high"]}<span class="arrow">→</span>{kpis["residual_high"]}',
               "High or critical", "inherent → residual"),
        figure(f'{kpis["reduction"]:.0%}', "Score reduction", "if all treatment lands"),
    ])

    overdue_rows = "".join(
        f"<tr><td class='id'>{r.id}</td><td>{esc(r.title)}<div class='sub'>{esc(r.remediation)}</div></td>"
        f"<td>{esc(short_owner(r.owner))}</td><td>{esc(r.status)}</td>"
        f"<td class='num'>{r.due:%d %b %Y}</td><td class='num late'>{days_overdue(r, today)} days</td></tr>"
        for r in kpis["overdue"]
    ) or "<tr><td colspan='6'>Nothing is overdue.</td></tr>"

    def chip(score):
        name = rating(score, bands)
        return f"<span class='chip' style='--c:{RATING_COLOURS[name]}'>{score} {name}</span>"

    register_rows = "".join(
        f"<tr><td class='id'>{r.id}</td><td>{esc(r.title)}</td><td>{esc(short_owner(r.owner))}</td>"
        f"<td>{chip(r.inherent)}</td><td>{chip(r.residual)}</td><td>{esc(r.status)}</td>"
        f"<td class='num'>{r.due:%d %b %Y}</td></tr>"
        for r in sorted(risks, key=lambda r: -r.inherent)
    )

    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Risk Dashboard</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="{GOOGLE_FONTS_CSS}" rel="stylesheet">
<style>
  :root {{ --paper: {PAPER}; --ink: {INK}; --muted: {MUTED}; --rule: {RULE}; --alert: {RATING_COLOURS['Critical']}; }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; background: var(--paper); color: var(--ink);
         font: 15px/1.55 '{SANS}', system-ui, sans-serif; }}
  main {{ max-width: 1120px; margin: 0 auto; padding: 48px 32px 64px; }}
  .kicker {{ font-size: 12px; letter-spacing: .12em; text-transform: uppercase; color: var(--muted);
            display: flex; justify-content: space-between; gap: 16px; flex-wrap: wrap; }}
  h1 {{ font: 600 44px/1.1 '{SERIF}', Georgia, serif; margin: 18px 0 14px; letter-spacing: -.01em; }}
  .lede {{ font: 400 19px/1.5 '{SERIF}', Georgia, serif; max-width: 760px; margin: 0; color: #3a3833; }}
  .masthead {{ border-bottom: 3px double var(--ink); padding-bottom: 28px; }}
  .figures {{ display: grid; grid-template-columns: repeat(5, 1fr); border-bottom: 1px solid var(--rule); }}
  .figure {{ padding: 22px 18px 20px; border-left: 1px solid var(--rule); }}
  .figure:first-child {{ border-left: 0; padding-left: 0; }}
  .value {{ font: 500 40px/1 '{MONO}', monospace; letter-spacing: -.03em; }}
  .value.alert {{ color: var(--alert); }}
  .arrow {{ color: var(--muted); font-size: 24px; margin: 0 6px; vertical-align: 6px; }}
  .label {{ font-weight: 600; font-size: 13px; margin-top: 10px; }}
  .note {{ font-size: 12px; color: var(--muted); }}
  section {{ margin-top: 52px; }}
  h2 {{ font: 600 24px/1.2 '{SERIF}', Georgia, serif; margin: 0 0 4px; }}
  h2 .num {{ font-family: '{MONO}', monospace; font-size: 14px; color: var(--muted); font-weight: 400;
             margin-right: 10px; vertical-align: 3px; }}
  .intro {{ color: var(--muted); margin: 0 0 18px; max-width: 720px; font-size: 14px; }}
  .chart {{ overflow-x: auto; }}
  svg {{ width: 100%; min-width: 720px; height: auto; display: block; }}
  svg g[id^="tip-"] {{ cursor: help; }}
  svg g[id^="tip-"]:hover path {{ stroke: var(--ink); }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13.5px; }}
  th {{ text-align: left; font-weight: 600; font-size: 11px; letter-spacing: .08em; text-transform: uppercase;
        color: var(--muted); border-bottom: 1px solid var(--ink); padding: 8px 10px 8px 0; }}
  td {{ border-bottom: 1px solid var(--rule); padding: 10px 10px 10px 0; vertical-align: top; }}
  td.id {{ font-family: '{MONO}', monospace; font-weight: 500; white-space: nowrap; }}
  td.num {{ font-family: '{MONO}', monospace; font-size: 12.5px; white-space: nowrap; }}
  td.late {{ color: var(--alert); font-weight: 500; }}
  .sub {{ color: var(--muted); font-size: 12.5px; margin-top: 2px; }}
  .chip {{ font-family: '{MONO}', monospace; font-size: 12px; white-space: nowrap; }}
  .chip::before {{ content: ""; display: inline-block; width: 9px; height: 9px; border-radius: 2px;
                   background: var(--c); margin-right: 6px; }}
  details summary {{ cursor: pointer; font-weight: 600; font-size: 14px; margin-bottom: 12px; }}
  .table-wrap {{ overflow-x: auto; }}
  footer {{ margin-top: 56px; padding-top: 16px; border-top: 1px solid var(--rule); font-size: 12px;
            color: var(--muted); display: flex; justify-content: space-between; gap: 16px; flex-wrap: wrap; }}
  @media (max-width: 760px) {{
    main {{ padding: 28px 16px 48px; }}
    h1 {{ font-size: 32px; }}
    .lede {{ font-size: 17px; }}
    .figures {{ grid-template-columns: repeat(2, 1fr); }}
    .figure, .figure:first-child {{ border-left: 0; padding-left: 0; border-bottom: 1px solid var(--rule); }}
    .value {{ font-size: 32px; }}
  }}
  @media print {{ main {{ padding: 0; }} section {{ break-inside: avoid; }} }}
</style>
</head>
<body>
<main>
  <header class="masthead">
    <div class="kicker"><span>Information security risk report · Saarland University</span>
      <span>ISO/IEC 27001:2022</span></div>
    <h1>Risk position on {today:%-d %B %Y}</h1>
    <p class="lede">{lede}</p>
  </header>

  <div class="figures">{figures}</div>

  <section>
    <h2><span class="num">01</span>Where the risks sit</h2>
    <p class="intro">Each risk placed by likelihood and impact, before and after the treatment the owners
      have planned. Hover a cell to see the risk titles.</p>
    <div class="chart">{svgs['heatmap']}</div>
  </section>

  <section>
    <h2><span class="num">02</span>What treatment buys</h2>
    <p class="intro">Inherent score against residual score for every risk. Anything right of the dashed line
      sits above the risk appetite and needs a treatment plan or a signed acceptance.</p>
    <div class="chart">{svgs['movement']}</div>
  </section>

  <section>
    <h2><span class="num">03</span>Who owns what</h2>
    <p class="intro">Risks per owner, split by remediation status.</p>
    <div class="chart">{svgs['owners']}</div>
  </section>

  <section>
    <h2><span class="num">04</span>Overdue remediations</h2>
    <p class="intro">Open risks whose remediation due date has passed, longest overdue first.</p>
    <div class="table-wrap"><table>
      <thead><tr><th>ID</th><th>Risk and planned remediation</th><th>Owner</th><th>Status</th>
        <th>Due</th><th>Overdue</th></tr></thead>
      <tbody>{overdue_rows}</tbody>
    </table></div>
  </section>

  <section>
    <details>
      <summary>All {len(risks)} risks as a table</summary>
      <div class="table-wrap"><table>
        <thead><tr><th>ID</th><th>Risk</th><th>Owner</th><th>Inherent</th><th>Residual</th>
          <th>Status</th><th>Due</th></tr></thead>
        <tbody>{register_rows}</tbody>
      </table></div>
    </details>
  </section>

  <footer>
    <span>Source: {esc(source_name)} · generated {date.today():%d %b %Y}</span>
    <span>Portfolio exercise. Not an official document of Saarland University.</span>
  </footer>
</main>
</body>
</html>
"""


def main():
    parser = argparse.ArgumentParser(description="Build the risk heatmap and KPI dashboard.")
    parser.add_argument("register", type=Path, help="risk register workbook (.xlsx)")
    parser.add_argument("--out", type=Path, default=Path(__file__).with_name("output"),
                        help="output folder (default: output/ next to this script)")
    parser.add_argument("--as-of", type=date.fromisoformat, default=date.today(),
                        help="reporting date for overdue checks, YYYY-MM-DD (default: today)")
    args = parser.parse_args()

    risks, bands, likelihood_names, impact_names = read_register(args.register)
    kpis = compute_kpis(risks, bands, args.as_of)
    load_fonts()
    args.out.mkdir(parents=True, exist_ok=True)

    tips = {}
    heatmap = draw_heatmap(risks, bands, likelihood_names, impact_names, tips)
    svgs = {
        "heatmap": save(heatmap, args.out, "heatmap", tips),
        "movement": save(draw_movement(risks, bands, kpis["appetite"]), args.out, "movement"),
        "owners": save(draw_owners(risks), args.out, "owners"),
    }
    page = build_page(risks, bands, kpis, args.as_of, args.register.name, svgs)
    (args.out / "risk-dashboard.html").write_text(page, encoding="utf-8")

    print(f"{kpis['total']} risks, {kpis['open']} open, {len(kpis['overdue'])} overdue as of {args.as_of}")
    print(f"Wrote heatmap.png, movement.png, owners.png and risk-dashboard.html to {args.out}")


if __name__ == "__main__":
    main()
